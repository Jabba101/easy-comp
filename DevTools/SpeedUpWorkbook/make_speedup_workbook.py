"""Builds the Easy Plugins speed-up workbook (results of the optimisation pass, 20 Oct 2026)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = "/home/claude/deliverables/Easy Plugins Speed-Up.xlsx"

FONT = "Arial"
INK = "1F2A33"
HEADER_FILL = PatternFill("solid", fgColor="2E4057")
SUBTLE_FILL = PatternFill("solid", fgColor="EEF1F4")
INPUT_FILL = PatternFill("solid", fgColor="FFF2B3")
LINE = Side(style="thin", color="C9D1D9")

def font(**kw):
    base = dict(name=FONT, size=10, color=INK)
    base.update(kw)
    return Font(**base)

MEASURED = font(color="0000FF")          # blue: typed-in measurements and assumptions
FORMULA = font()                          # black: worked out from other cells
LINKED = font(color="008000")             # green: taken from another sheet

def title(ws, text, subtitle):
    ws["A1"] = text
    ws["A1"].font = font(size=14, bold=True)
    ws["A2"] = subtitle
    ws["A2"].font = font(size=10, italic=True, color="52606D")

def header(ws, row, labels):
    for col, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=col, value=label)
        c.font = font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(bottom=LINE)
    ws.row_dimensions[row].height = 44

def body_cell(ws, row, col, value, style=FORMULA, fmt=None, wrap=False, fill=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = style
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical="top")
    c.border = Border(bottom=LINE)
    if fill:
        c.fill = fill
    return c

def lines_needed(text, width_chars):
    return max(1, -(-len(text) // max(1, int(width_chars * 1.05))))

def notes(ws, row, lines, width_cols):
    ws.cell(row=row, column=1, value="Notes").font = font(bold=True)
    total_width = sum(ws.column_dimensions[get_column_letter(c)].width for c in range(1, width_cols + 1))
    for i, text in enumerate(lines, start=1):
        ws.merge_cells(start_row=row + i, start_column=1, end_row=row + i, end_column=width_cols)
        c = ws.cell(row=row + i, column=1, value=text)
        c.font = font(color="3E4C59")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row + i].height = 13.5 * lines_needed(text, total_width) + 3

def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

wb = Workbook()

# ---------------------------------------------------------------------------
# 1. CPU use
# ---------------------------------------------------------------------------
ws = wb.active
ws.title = "CPU use"
title(ws, "Processor use, before and after the speed-up",
      "Share of one processor core each plugin needs while playing (lower is better).")

header(ws, 4, ["Plugin", "Preset", "AUTO-GAIN", "Audio", "Before\n(share of one core)", "After\n(share of one core)",
               "Times less", "Processor time saved"])

rows = [
    # plugin, preset, auto-gain, audio, before, after  (% of one core, best of five)
    ("Easy Comp", "Drum Bus - Glue", "Off", "Rock drum bus", 0.415, 0.136),
    ("Easy Comp", "Drum Bus - Glue", "On", "Rock drum bus", 0.377, 0.136),
    ("Easy Comp", "Vocal - Pop & R&B", "Off", "Rock drum bus", 0.419, 0.137),
    ("Easy Comp", "Vocal - Pop & R&B", "On", "Rock drum bus", 0.419, 0.137),
    ("Easy Comp", "Drum Bus - Glue", "On", "Digital silence", 0.105, 0.004),
    ("Easy Transient", "Drum Bus - Punchier", "Off", "Rock drum bus", 0.605, 0.171),
    ("Easy Transient", "Drum Bus - Punchier", "On", "Rock drum bus", 0.631, 0.170),
    ("Easy Transient", "Drum Bus - Punchier", "On", "Digital silence", 0.153, 0.008),
]

widths(ws, [16, 22, 12, 18, 16, 16, 12, 14])
first = 5
for i, (plugin, preset, auto, audio, before, after) in enumerate(rows):
    r = first + i
    body_cell(ws, r, 1, plugin)
    body_cell(ws, r, 2, preset)
    body_cell(ws, r, 3, auto)
    body_cell(ws, r, 4, audio)
    body_cell(ws, r, 5, before / 100.0, MEASURED, "0.000%")
    body_cell(ws, r, 6, after / 100.0, MEASURED, "0.000%")
    body_cell(ws, r, 7, f'=IF(F{r}>0,E{r}/F{r},"")', FORMULA, '0.0"×"')
    body_cell(ws, r, 8, f'=IF(E{r}>0,(E{r}-F{r})/E{r},"")', FORMULA, "0%")
last = first + len(rows) - 1

notes(ws, last + 2, [
    "Blue figures are measurements; black ones are worked out from them, so they update if you change a measurement.",
    "Measured on 20 October 2026 on Claude's cloud test computer (Intel Xeon, 2.1 GHz): 120 seconds of stereo audio at 48 kHz, "
    "in blocks of 256 samples, best of five runs. Your i7-2600K should be in the same range, possibly up to 1.5× these figures.",
    "Before: Easy Comp 1.4.0 and Easy Transient 1.1.1. After: Easy Comp 1.5.0 and Easy Transient 1.2.0.",
    "\"Rock drum bus\" is the test recording the drum presets were tuned on; the vocal preset was timed on it too "
    "(a compressor's processor use hardly depends on the audio, only on whether there is any). "
    "\"Digital silence\" is a track with nothing on it, or the DAW stopped.",
    "Resting in silence only applies to true digital silence: a track with any sound at all, such as mic bleed, is processed as "
    "normal (at the lower cost shown in the other rows).",
], 8)

ws.freeze_panes = "A5"

# ---------------------------------------------------------------------------
# 2. Sound check
# ---------------------------------------------------------------------------
ws = wb.create_sheet("Sound check")
title(ws, "Does it still sound the same?",
      "The old and new versions were run side by side on the same audio, and the difference between them measured.")

widths(ws, [30, 52, 11, 16, 14, 12])
ws["A4"] = "Pass mark"
ws["A4"].font = font(bold=True)
marks = [
    ("The difference must be at least this far below the music (dB)", 100),
    ("No moment's level may change by more than (dB)", 0.001),
    ("For comparison: the smallest level change most people can hear (dB, approximate)", 0.2),
]
for i, (label, value) in enumerate(marks):
    r = 5 + i
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws.cell(row=r, column=1, value=label).font = font()
    c = ws.cell(row=r, column=5, value=value)
    c.font = MEASURED
    c.number_format = "0.000" if value < 1 else "0"
    c.alignment = Alignment(horizontal="left")

header(ws, 9, ["Plugin", "What was compared", "Situations", "Difference: how far below the music (dB)",
               "Largest level change (dB)", "Passes?"])

checks = [
    ("Easy Comp", "Every preset on the recordings it was tuned on, AUTO-GAIN off and on", 168, 109.6, 0.000047),
    ("Easy Comp", "Every preset, with gaps of digital silence (1 sample to 5 seconds) cut into the audio", 114, 110.3, 0.000048),
    ("Easy Comp", "Knob moves, preset loads, BYPASS and AUTO-GAIN switched during silence; random automation", 20, 118.5, 0.000042),
    ("Easy Comp", "Buffer sizes from 1 to 4096 samples; 44.1, 96 and 192 kHz; mono", 16, 118.6, 0.000071),
    ("Easy Comp", "Awkward signals (DC, square waves, very quiet noise, a hot sine) and silences up to a minute", 6, 104.8, 0.000241),
    ("Easy Transient", "Every preset on the recordings it was tuned on, AUTO-GAIN off and on", 120, 126.4, 0.000006),
    ("Easy Transient", "Every preset, with gaps of digital silence (1 sample to 5 seconds) cut into the audio", 76, 126.3, 0.000006),
    ("Easy Transient", "Knob moves, preset loads, BYPASS and AUTO-GAIN switched during silence; random automation", 16, 123.9, 0.000009),
    ("Easy Transient", "Buffer sizes from 1 to 4096 samples; 44.1, 96 and 192 kHz; mono", 10, 128.2, 0.000009),
    ("Easy Transient", "Awkward signals and silences up to a minute", 6, 124.5, 0.000011),
    ("Easy Comp", "The 25 reference recordings made with version 1.2, against the new version", 25, 111.4, 0.000087),
]
first = 10
for i, (plugin, what, n, below, level) in enumerate(checks):
    r = first + i
    body_cell(ws, r, 1, plugin)
    body_cell(ws, r, 2, what, wrap=True)
    body_cell(ws, r, 3, n, MEASURED, "0")
    body_cell(ws, r, 4, below, MEASURED, "0.0")
    body_cell(ws, r, 5, level, MEASURED, "0.000000")
    body_cell(ws, r, 6, f'=IF(AND(D{r}>=$E$5,E{r}<=$E$6),"Yes","No")').alignment = Alignment(horizontal="center", vertical="top")
last = first + len(checks) - 1
tot = last + 1
for col in range(1, 7):
    ws.cell(row=tot, column=col).fill = SUBTLE_FILL
body_cell(ws, tot, 1, "All together", font(bold=True), fill=SUBTLE_FILL)
body_cell(ws, tot, 2, "Worst case across everything", font(bold=True), fill=SUBTLE_FILL)
body_cell(ws, tot, 3, f"=SUM(C{first}:C{last})", font(bold=True), "0", fill=SUBTLE_FILL)
body_cell(ws, tot, 4, f"=MIN(D{first}:D{last})", font(bold=True), "0.0", fill=SUBTLE_FILL)
body_cell(ws, tot, 5, f"=MAX(E{first}:E{last})", font(bold=True), "0.000000", fill=SUBTLE_FILL)
body_cell(ws, tot, 6, f'=IF(COUNTIF(F{first}:F{last},"Yes")=ROWS(F{first}:F{last}),"All pass","Check")', font(bold=True), fill=SUBTLE_FILL).alignment = Alignment(horizontal="center", vertical="top")
body_cell(ws, tot + 1, 2, "Worst level change, as a share of the smallest audible change", font(italic=True))
body_cell(ws, tot + 1, 5, f"=E{tot}/$E$7", font(italic=True), "0.000%")

other = tot + 3
ws.cell(row=other, column=1, value="Other checks").font = font(bold=True)
header(ws, other + 1, ["Check", "Result", "", "", "", ""])
ws.merge_cells(start_row=other + 1, start_column=2, end_row=other + 1, end_column=6)
other_rows = [
    ("New engine vs. doing every step one sample at a time", "Identical, bit for bit, in all 552 situations above (the 25 reference recordings aside)"),
    ("Skipping through silence vs. working through it", "At least 121 dB apart, no level change over 0.00002 dB"),
    ("Automatic tests, Easy Comp", "289 of 289 passed"),
    ("Automatic tests, Easy Transient", "58 of 58 passed"),
    ("pluginval, strictest level (10)", "Both plugins passed"),
    ("Windows build on GitHub", "Succeeded: Easy Comp.vst3 and Easy Transient.vst3"),
]
for i, (check, result) in enumerate(other_rows):
    r = other + 2 + i
    body_cell(ws, r, 1, check, wrap=True)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    body_cell(ws, r, 2, result, wrap=True)
    ws.row_dimensions[r].height = 28

notes(ws, other + 2 + len(other_rows) + 1, [
    "\"How far below the music\" compares the largest difference between old and new with the loudest moment of the audio: "
    "100 dB below is a ratio of 1 to 100,000.",
    "Figures from the final comparison runs on 20 October 2026, on the code released as Easy Comp 1.5.0 and Easy Transient 1.2.0. "
    "Blue figures are measurements; the Passes? column and the totals are worked out from them and the pass mark above.",
    "The smallest audible change is an approximate, commonly quoted figure for level differences (around 0.2 dB in ideal "
    "listening conditions; most people need more).",
], 6)

for r in range(first, last + 1):
    ws.row_dimensions[r].height = 30
ws.freeze_panes = "A10"

# ---------------------------------------------------------------------------
# 3. What changed
# ---------------------------------------------------------------------------
ws = wb.create_sheet("What changed")
title(ws, "What changed in the engine",
      "Easy Comp 1.5.0 and Easy Transient 1.2.0, released 20 October 2026.")
header(ws, 4, ["Change", "What it does", "Effect on the sound"])
changes = [
    ("Quicker decibel maths",
     "The plugins turn volumes into decibels and back for every sample. They now use a shortcut accurate to about 0.00002 dB "
     "instead of the general-purpose version. It's exact at 0 dB, so a setting meant to change nothing still changes nothing.",
     "Differences at least 104 dB below the music"),
    ("Several samples at once",
     "The work goes through 64 samples at a time, and the parts that don't depend on the previous sample are worked out "
     "four samples at once, using SSE2 (which every 64-bit processor has). Left and right are measured together for AUTO-GAIN.",
     "None: identical, bit for bit, to doing it one sample at a time"),
    ("Resting in silence",
     "On a digitally silent track (empty stretches between regions, or the DAW stopped), the plugins skip ahead in one step "
     "once the last sound has died away, instead of working through every sample.",
     "Differences at least 121 dB below the music"),
    ("Only recalculate what moved",
     "Values worked out from the knobs are only worked out again when a knob moves, not on every block of audio.",
     "None"),
    ("Shared by every plugin",
     "All of this lives in the shared building blocks, so the gated reverb and the channel strip get it automatically.",
     "Not applicable"),
]
for i, (name, what, effect) in enumerate(changes):
    r = 5 + i
    body_cell(ws, r, 1, name, font(bold=True), wrap=True)
    body_cell(ws, r, 2, what, wrap=True)
    body_cell(ws, r, 3, effect, wrap=True)
    ws.row_dimensions[r].height = 58
widths(ws, [24, 70, 30])
ws.freeze_panes = "A5"

# ---------------------------------------------------------------------------
# 4. Channel strip estimate
# ---------------------------------------------------------------------------
ws = wb.create_sheet("Channel strip estimate")
title(ws, "Channel strip: a rough estimate",
      "What a strip might need if each module costs about what Easy Comp and Easy Transient do now.")
ws["A3"] = "Yellow cells are estimates you can change; the green ones come from the CPU use sheet."
ws["A3"].font = font(color="52606D")
widths(ws, [28, 18, 62])
header(ws, 5, ["Module", "Share of one core", "Where the figure comes from"])
modules = [
    ("Compressor", "='CPU use'!F6", LINKED, None, "Measured: Easy Comp 1.5.0 on the rock drum bus, AUTO-GAIN on"),
    ("Transient shaper", "='CPU use'!F11", LINKED, None, "Measured: Easy Transient 1.2.0 on the rock drum bus, AUTO-GAIN on"),
    ("Gate", 0.00136, MEASURED, INPUT_FILL, "Estimate, not built yet: assumed about the same as the compressor"),
    ("EQ", 0.00136, MEASURED, INPUT_FILL, "Estimate, not built yet: assumed about the same as the compressor"),
]
for i, (name, value, style, fill, source) in enumerate(modules):
    r = 6 + i
    body_cell(ws, r, 1, name)
    body_cell(ws, r, 2, value, style, "0.000%", fill=fill)
    body_cell(ws, r, 3, source, wrap=True).alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    ws.row_dimensions[r].height = 13.5 * lines_needed(source, 62) + 3
for col in range(1, 4):
    ws.cell(row=10, column=col).fill = SUBTLE_FILL
body_cell(ws, 10, 1, "Whole strip", font(bold=True), fill=SUBTLE_FILL)
body_cell(ws, 10, 2, "=SUM(B6:B9)", font(bold=True), "0.000%", fill=SUBTLE_FILL)
body_cell(ws, 10, 3, "Sum of the modules above", fill=SUBTLE_FILL).alignment = Alignment(vertical="top", indent=1)

body_cell(ws, 12, 1, "Allowance for an older PC")
body_cell(ws, 12, 2, 1.5, MEASURED, '0.0"×"', fill=INPUT_FILL)
body_cell(ws, 12, 3, "Estimate: the i7-2600K might be up to this much slower than the test computer", wrap=True).alignment = Alignment(wrap_text=True, vertical="top", indent=1)
ws.row_dimensions[12].height = 13.5 * lines_needed("Estimate: the i7-2600K might be up to this much slower than the test computer", 62) + 3
body_cell(ws, 13, 1, "Whole strip on an older PC", font(bold=True))
body_cell(ws, 13, 2, "=B10*B12", font(bold=True), "0.000%")
body_cell(ws, 13, 3, "Whole strip × allowance").alignment = Alignment(vertical="top", indent=1)

notes(ws, 15, [
    "A rough planning figure only. The gate and EQ don't exist yet, and an analog-style drive stage would need oversampling, "
    "which usually makes it the most expensive part of a strip.",
], 3)

# Printing: landscape, each sheet one page wide
for sheet in wb.worksheets:
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_view.showGridLines = False

# Arial for anything left at the default style, too
wb._named_styles["Normal"].font = Font(name=FONT, size=10)

wb.save(OUT)
print("saved", OUT)
